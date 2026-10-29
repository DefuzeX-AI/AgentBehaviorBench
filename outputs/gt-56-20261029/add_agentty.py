"""Add the restored agentty #116 to the existing GT workbook."""

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font


repo = Path(__file__).resolve().parents[2]
source = repo / "resources/agents/93-agentty/ground_truth/issue_116/verify.case.json"
issue_dir = source.parent
assert any((issue_dir / "runs").glob("*/trajectory/*/trajectory.json"))
definition = json.loads(source.read_text(encoding="utf-8"))
locations = definition["fault location"]
cause = definition["Root cause"]
assert locations and cause.strip()
location_lines = []
for location in locations:
    start = location.get("line_start")
    end = location.get("line_end")
    line = f":{start}" if start else ""
    if end and end != start:
        line += f"-{end}"
    function = f" ({location['function']})" if location.get("function") else ""
    location_lines.append(f"{location['file']}{line}{function}")

old_path = Path(__file__).with_name("gt-56.xlsx")
new_path = Path(__file__).with_name("gt-57.xlsx")
assert old_path.is_file()
workbook = load_workbook(old_path)
sheet = workbook["GT cases"]
assert sheet.max_row == 57 and sheet.cell(2, 1).value == "article-explainer"
sheet.insert_rows(2)
for column in range(1, 7):
    template = sheet.cell(3, column)
    cell = sheet.cell(2, column)
    cell._style = copy(template._style)
    cell.alignment = copy(template.alignment)
    cell.protection = copy(template.protection)

repository = definition["issue"]["repository"]
issue_url = definition["issue"]["url"]
values = ["agentty", repository, definition["issue"]["number"], issue_url,
          "\n".join(location_lines), cause]
for column, value in enumerate(values, 1):
    sheet.cell(2, column).value = value
sheet.cell(2, 2).hyperlink = f"https://github.com/{repository}"
sheet.cell(2, 4).hyperlink = issue_url
for column in (2, 4):
    sheet.cell(2, column).font = Font(name="Arial", size=10, color="1459A5", underline="single")
sheet.row_dimensions[2].height = min(320, max(34,
    ((len(values[4]) + 83) // 84) * 15,
    ((len(cause) + 107) // 108) * 15))
for row in range(2, 59):
    repo_cell = sheet.cell(row, 2)
    issue_cell = sheet.cell(row, 4)
    repo_cell.hyperlink = f"https://github.com/{repo_cell.value}"
    issue_cell.hyperlink = issue_cell.value
    for cell in (repo_cell, issue_cell):
        cell.font = Font(name="Arial", size=10, color="1459A5", underline="single")
table = sheet.tables["GroundTruthCases"]
table.ref = "A1:F58"
if table.autoFilter:
    table.autoFilter.ref = table.ref
workbook.save(new_path)

check = load_workbook(new_path)
sheet = check["GT cases"]
assert sheet.max_row == 58 and sheet.tables["GroundTruthCases"].ref == "A1:F58"
assert sheet.freeze_panes == "C2"
assert sum(sheet.cell(row, 1).value == "agentty" and sheet.cell(row, 3).value == 116
           for row in range(2, 59)) == 1
assert all(sheet.cell(row, 2).hyperlink and sheet.cell(row, 4).hyperlink
           for row in range(2, 59))
assert all(sheet.cell(row, 5).value and sheet.cell(row, 6).value
           for row in range(2, 59))
assert sheet.cell(2, 6).value == cause
print("Verified: 57 GT rows, 114 native hyperlinks, agentty #116 restored")
print(new_path)
