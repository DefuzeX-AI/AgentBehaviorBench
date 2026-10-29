"""Add native Excel hyperlinks because the artifact engine cannot calculate HYPERLINK."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font


path = Path(__file__).with_name("gt-56.xlsx")
workbook = load_workbook(path)
sheet = workbook["GT cases"]
assert sheet.max_row == 57
for row in range(2, sheet.max_row + 1):
    repository = sheet.cell(row, 2)
    issue_link = sheet.cell(row, 4)
    assert isinstance(repository.value, str) and "/" in repository.value
    assert isinstance(issue_link.value, str) and issue_link.value.startswith("https://")
    repository.hyperlink = f"https://github.com/{repository.value}"
    issue_link.hyperlink = issue_link.value
    for cell in (repository, issue_link):
        cell.font = Font(name="Arial", size=10, color="1459A5", underline="single")

temporary = path.with_name("gt-56.links.tmp.xlsx")
workbook.save(temporary)
temporary.replace(path)

verified = load_workbook(path, read_only=False)
sheet = verified["GT cases"]
assert sheet.max_row == 57
assert len(sheet.tables) == 1
assert all(sheet.cell(row, 2).hyperlink and sheet.cell(row, 4).hyperlink
           for row in range(2, 58))
assert all(sheet.cell(row, 5).value and sheet.cell(row, 6).value
           for row in range(2, 58))
print("Verified: 56 rows, 112 native hyperlinks, all locations and root causes present")
